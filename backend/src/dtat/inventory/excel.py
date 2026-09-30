"""Excel template, export and import of the inventory.

Import rules (also written into the «Инструкция» sheet):
- objects are matched by natural keys: site code, eNB ID + Cell ID, board number, IMEI or name;
- only columns present in the sheet are applied; an empty cell clears an optional value;
- nothing is deleted;
- the import is all-or-nothing: any row error rolls back the whole file.
"""

from collections.abc import Callable, Iterator
from dataclasses import dataclass
from datetime import datetime
from enum import StrEnum
from io import BytesIO
from typing import Any, Literal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel, ValidationError
from sqlalchemy.orm import Session

from dtat.audit.service import Actor
from dtat.errors import DomainError, InvalidDataError
from dtat.inventory import service
from dtat.inventory.labels import LABELS, label, parse_enum
from dtat.inventory.lte import make_eci
from dtat.inventory.models import AssetKind, DeviceKind, Site, SiteKind, Status
from dtat.inventory.schemas import (
    AssetCreate,
    AssetUpdate,
    CellCreate,
    CellUpdate,
    DeviceCreate,
    DeviceUpdate,
    ENodeBCreate,
    ENodeBUpdate,
    SiteCreate,
    SiteUpdate,
)

Kind = Literal["str", "int", "float", "enum"]
Outcome = Literal["created", "updated", "unchanged"]


@dataclass(frozen=True)
class Column:
    key: str
    header: str
    kind: Kind = "str"
    required: bool = False
    enum: type[StrEnum] | None = None
    width: int = 14
    export_only: bool = False


@dataclass(frozen=True)
class SheetSpec:
    title: str
    columns: tuple[Column, ...]

    def column(self, key: str) -> Column:
        return next(c for c in self.columns if c.key == key)


NOTES = Column("notes", "Примечание", width=30)
STATUS = Column("status", "Статус", "enum", enum=Status)

ASSETS = SheetSpec(
    "Техника",
    (
        Column("name", "Бортовой номер", required=True, width=18),
        Column("kind", "Тип", "enum", enum=AssetKind, width=18),
        Column("model", "Модель", width=20),
        NOTES,
    ),
)
SITES = SheetSpec(
    "Сайты",
    (
        Column("code", "Код сайта", required=True),
        Column("name", "Название", width=24),
        Column("kind", "Тип площадки", "enum", enum=SiteKind, width=16),
        STATUS,
        Column("lat", "Широта", "float", required=True, width=12),
        Column("lon", "Долгота", "float", required=True, width=12),
        Column("structure_type", "Тип опоры"),
        Column("structure_height_m", "Высота опоры, м", "float"),
        Column("ground_elevation_m", "Отметка земли, м", "float"),
        NOTES,
    ),
)
CELLS = SheetSpec(
    "Соты",
    (
        Column("site_code", "Код сайта", required=True),
        Column("enb_id", "eNB ID", "int", required=True, width=10),
        Column("enb_name", "Имя eNB", width=16),
        Column("enb_site_code", "Код сайта eNB", width=14),
        Column("local_cell_id", "Cell ID", "int", required=True, width=9),
        Column("eci", "ECI (расчётный)", "int", export_only=True, width=16),
        Column("name", "Имя соты", width=18),
        STATUS,
        Column("pci", "PCI", "int", required=True, width=8),
        Column("earfcn_dl", "EARFCN DL", "int", required=True, width=11),
        Column("earfcn_ul", "EARFCN UL", "int", width=11),
        Column("bandwidth_mhz", "Полоса, МГц", "float", width=12),
        Column("tac", "TAC", "int", width=8),
        Column("max_tx_power_dbm", "Мощность, дБм", "float"),
        Column("antenna_model", "Модель антенны", width=18),
        Column("height_m", "Высота подвеса, м", "float"),
        Column("azimuth_deg", "Азимут, °", "float", width=10),
        Column("mech_tilt_deg", "Мех. тилт, °", "float", width=11),
        Column("elec_tilt_deg", "Эл. тилт, °", "float", width=11),
        Column("beamwidth_deg", "Ширина ДН, °", "float", width=12),
        NOTES,
    ),
)
DEVICES = SheetSpec(
    "Устройства",
    (
        Column("name", "Название", required=True, width=20),
        Column("kind", "Тип", "enum", required=True, enum=DeviceKind),
        STATUS,
        Column("imei", "IMEI", width=18),
        Column("imsi", "IMSI", width=18),
        Column("iccid", "ICCID", width=22),
        Column("model", "Модель", width=18),
        Column("asset_name", "Техника (бортовой №)", width=20),
        Column("role", "Роль", width=18),
        NOTES,
    ),
)

# Import order matters: sites before cells, assets before devices.
SHEETS = (SITES, CELLS, ASSETS, DEVICES)
IMPORT_ORDER = (ASSETS, SITES, CELLS, DEVICES)

# Non-nullable fields that have a default: an empty cell means "keep as is / use default".
_KEEP_IF_EMPTY = {"status", "kind"}

MAX_ROWS_WITH_VALIDATION = 5000


# --- report ----------------------------------------------------------------------------------


class RowError(BaseModel):
    row: int
    message: str


class SheetReport(BaseModel):
    sheet: str
    present: bool = True
    created: int = 0
    updated: int = 0
    unchanged: int = 0
    errors: list[RowError] = []


class ImportReport(BaseModel):
    dry_run: bool
    applied: bool
    sheets: list[SheetReport]
    total_errors: int


# --- writing ---------------------------------------------------------------------------------

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_REQUIRED_FILL = PatternFill("solid", fgColor="C00000")
_EXPORT_ONLY_FILL = PatternFill("solid", fgColor="7F7F7F")

INSTRUCTIONS = [
    "Шаблон инвентаря частной LTE-сети",
    "",
    "Листы: «Сайты», «Соты», «Техника», «Устройства».",
    "Названия листов и заголовки столбцов не менять.",
    "Красные заголовки: обязательные столбцы. Серые: только для чтения, при загрузке игнорируются.",
    "",
    "Как загрузка сопоставляет строки с объектами в системе:",
    "  • сайт: по коду сайта;",
    "  • сота: по паре eNB ID + Cell ID (из них рассчитывается ECI = eNB ID × 256 + Cell ID);",
    "    «Код сайта» соты — где установлена антенна. «Код сайта eNB» заполняется только для",
    "    выносных секторов и DAS в другом здании: там указывается сайт, где стоит сам eNodeB.",
    "  • техника: по бортовому номеру;",
    "  • устройство: по IMEI, если указан, иначе по названию.",
    "Найденные объекты обновляются, новые создаются. Удаления через файл не выполняются.",
    "Применяются только столбцы, присутствующие на листе.",
    "Пустая ячейка очищает необязательное поле; пустые «Статус» и «Тип» не меняют значение.",
    "Если в файле есть хотя бы одна ошибка, не применяется ничего:",
    "исправьте ошибки и загрузите файл снова.",
    "",
    "Координаты: WGS-84, десятичные градусы (55.123456). Азимут: 0–359.9°, от севера по часовой.",
    "eNB ID сайта на Ericsson: параметр ENodeBFunction.eNBId; Cell ID: EUtranCellFDD.cellId.",
    "PCI = 3 × physicalLayerCellIdGroup + physicalLayerSubCellId.",
    "IMEI: 15 цифр. ICCID храните как текст, иначе Excel исказит длинное число.",
    "",
    "Допустимые значения:",
    *(
        f"  • {title}: " + ", ".join(LABELS[enum_cls].values())
        for title, enum_cls in (
            ("Статус", Status),
            ("Тип площадки", SiteKind),
            ("Тип техники", AssetKind),
            ("Тип устройства", DeviceKind),
        )
    ),
]


def _write_sheet(wb: Workbook, spec: SheetSpec, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(spec.title)
    for col_idx, column in enumerate(spec.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=column.header)
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.fill = (
            _EXPORT_ONLY_FILL
            if column.export_only
            else _REQUIRED_FILL
            if column.required
            else _HEADER_FILL
        )
        letter = get_column_letter(col_idx)
        ws.column_dimensions[letter].width = column.width
        if column.kind == "str":
            # Text format keeps IMEI/ICCID and codes like "007" intact, also in rows typed later.
            ws.column_dimensions[letter].number_format = "@"
        if column.enum is not None:
            options = ",".join(LABELS[column.enum].values())
            validation = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
            validation.add(f"{letter}2:{letter}{MAX_ROWS_WITH_VALIDATION}")
            ws.add_data_validation(validation)

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, column in enumerate(spec.columns, start=1):
            value = row.get(column.key)
            if isinstance(value, StrEnum):
                value = label(value)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if column.kind == "str":
                cell.number_format = "@"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(spec.columns))}{max(len(rows) + 1, 1)}"


def _export_rows(session: Session) -> dict[str, list[dict[str, Any]]]:
    sites = [
        {c.key: getattr(site, c.key) for c in SITES.columns} for site in service.list_sites(session)
    ]
    cells = []
    for cell in service.list_cells(session):
        row = {c.key: getattr(cell, c.key, None) for c in CELLS.columns}
        row.update(
            site_code=cell.site.code,
            enb_id=cell.enodeb.enb_id,
            enb_name=cell.enodeb.name,
            # Filled only for remote sectors, whose eNodeB is on another site.
            enb_site_code=cell.enodeb.site.code if cell.enodeb.site_id != cell.site_id else None,
        )
        cells.append(row)
    assets = [
        {c.key: getattr(asset, c.key) for c in ASSETS.columns}
        for asset in service.list_assets(session)
    ]
    devices = []
    for device in service.list_devices(session):
        row = {c.key: getattr(device, c.key, None) for c in DEVICES.columns}
        row["asset_name"] = device.asset.name if device.asset else None
        devices.append(row)
    return {SITES.title: sites, CELLS.title: cells, ASSETS.title: assets, DEVICES.title: devices}


def build_workbook(session: Session | None) -> bytes:
    """Export of the whole inventory, or an empty template when `session` is None."""
    rows = _export_rows(session) if session is not None else {}
    wb = Workbook()
    info = wb.active
    assert info is not None
    info.title = "Инструкция"
    info.column_dimensions["A"].width = 110
    for i, line in enumerate(INSTRUCTIONS, start=1):
        info.cell(row=i, column=1, value=line).font = Font(bold=(i == 1), size=12 if i == 1 else 11)
    for spec in SHEETS:
        _write_sheet(wb, spec, rows.get(spec.title, []))
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --- reading ---------------------------------------------------------------------------------


def _normalize_header(text: Any) -> str:
    return str(text or "").replace("*", "").strip().lower()


def _to_text(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _convert(column: Column, value: Any) -> Any:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if column.kind == "str":
        return _to_text(value)
    if column.kind == "enum":
        assert column.enum is not None
        parsed = parse_enum(column.enum, _to_text(value))
        if parsed is None:
            allowed = ", ".join(LABELS[column.enum].values())
            raise InvalidDataError(f"{column.header}: «{value}» не из списка ({allowed})")
        return parsed
    if isinstance(value, datetime):
        raise InvalidDataError(f"{column.header}: ожидается число, а не дата")
    try:
        number = float(str(value).replace(",", ".").replace(" ", ""))
    except ValueError:
        raise InvalidDataError(f"{column.header}: «{value}» не является числом") from None
    if column.kind == "int":
        if not number.is_integer():
            raise InvalidDataError(f"{column.header}: ожидается целое число, получено {value}")
        return int(number)
    return number


def _iter_rows(ws: Worksheet, spec: SheetSpec) -> Iterator[tuple[int, dict[str, Any]]]:
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None) or ()
    by_header = {_normalize_header(c.header): c for c in spec.columns if not c.export_only}
    positions: dict[int, Column] = {}
    for idx, text in enumerate(header):
        column = by_header.get(_normalize_header(text))
        if column is not None:
            positions[idx] = column
    present = {c.key for c in positions.values()}
    missing = [c.header for c in spec.columns if c.required and c.key not in present]
    if missing:
        raise InvalidDataError("Нет обязательных столбцов: " + ", ".join(missing))

    for row_num, values in enumerate(rows, start=2):
        if values is None or all(v is None or str(v).strip() == "" for v in values):
            continue
        raw = {
            col.key: (values[idx] if idx < len(values) else None) for idx, col in positions.items()
        }
        yield row_num, raw


def _pydantic_message(spec: SheetSpec, exc: ValidationError) -> str:
    parts = []
    for error in exc.errors():
        key = str(error["loc"][0]) if error["loc"] else ""
        header = next((c.header for c in spec.columns if c.key == key), key)
        message = str(error["msg"]).removeprefix("Value error, ")
        parts.append(f"{header}: {message}" if header else message)
    return "; ".join(parts)


def _values(spec: SheetSpec, raw: dict[str, Any]) -> dict[str, Any]:
    values = {key: _convert(spec.column(key), value) for key, value in raw.items()}
    missing = [
        spec.column(k).header for k, v in values.items() if spec.column(k).required and v is None
    ]
    if missing:
        raise InvalidDataError("Не заполнено: " + ", ".join(missing))
    for key in _KEEP_IF_EMPTY:
        if values.get(key, ...) is None:
            del values[key]
    return values


def _outcome(changed: bool) -> Outcome:
    return "updated" if changed else "unchanged"


def _import_asset(
    session: Session, actor: Actor, raw: dict[str, Any], _at: datetime | None
) -> Outcome:
    values = _values(ASSETS, raw)
    asset = service.find_asset_by_name(session, values["name"])
    if asset is None:
        service.create_asset(session, actor, AssetCreate.model_validate(values))
        return "created"
    _, changed = service.update_asset(session, actor, asset.id, AssetUpdate.model_validate(values))
    return _outcome(changed)


def _import_site(
    session: Session, actor: Actor, raw: dict[str, Any], at: datetime | None
) -> Outcome:
    values = _values(SITES, raw)
    site = service.find_site_by_code(session, values["code"])
    if site is None:
        service.create_site(session, actor, SiteCreate.model_validate(values))
        return "created"
    update = SiteUpdate.model_validate({**values, "effective_at": at})
    _, changed = service.update_site(session, actor, site.id, update)
    return _outcome(changed)


def _import_cell(
    session: Session, actor: Actor, raw: dict[str, Any], at: datetime | None
) -> Outcome:
    values = _values(CELLS, raw)
    site_code = values.pop("site_code")
    enb_id = values.pop("enb_id")
    enb_name = values.pop("enb_name", None)
    enb_site_code = values.pop("enb_site_code", None)
    local_cell_id = values.pop("local_cell_id")

    def site_or_error(code: str) -> Site:
        site = service.find_site_by_code(session, code)
        if site is None:
            raise InvalidDataError(f"Сайт {code} не найден: добавьте его на лист «Сайты»")
        return site

    # "Код сайта" is where the antenna is; "Код сайта eNB" is set only for remote sectors.
    site = site_or_error(site_code)
    enb_site = site_or_error(enb_site_code) if enb_site_code else site
    enodeb = service.find_enodeb_by_enb_id(session, enb_id)
    renumbered = False
    if enodeb is None:
        # A known cell under another eNB ID: the eNB ID was corrected (the operator's numbering
        # became known). Renumber that eNodeB instead of creating a duplicate; ECIs follow.
        known = service.find_cell_by_name(session, values["name"]) if values.get("name") else None
        if known is not None and known.local_cell_id == local_cell_id:
            enodeb, _ = service.update_enodeb(
                session,
                actor,
                known.enodeb_id,
                ENodeBUpdate.model_validate({"enb_id": enb_id, "effective_at": at}),
            )
            renumbered = True
    if enodeb is None:
        enodeb = service.create_enodeb(
            session, actor, ENodeBCreate(site_id=enb_site.id, enb_id=enb_id, name=enb_name)
        )
    else:
        enodeb_changes: dict[str, Any] = {"effective_at": at}
        if enb_site_code and enodeb.site_id != enb_site.id:
            enodeb_changes["site_id"] = enb_site.id
        if enb_name is not None:
            enodeb_changes["name"] = enb_name
        service.update_enodeb(
            session, actor, enodeb.id, ENodeBUpdate.model_validate(enodeb_changes)
        )

    cell = service.find_cell_by_eci(session, make_eci(enb_id, local_cell_id))
    if cell is None:
        data = CellCreate.model_validate(
            {**values, "enodeb_id": enodeb.id, "site_id": site.id, "local_cell_id": local_cell_id}
        )
        service.create_cell(session, actor, data)
        return "created"
    _, changed = service.update_cell(
        session,
        actor,
        cell.id,
        CellUpdate.model_validate({**values, "site_id": site.id, "effective_at": at}),
    )
    return _outcome(changed or renumbered)


def _import_device(
    session: Session, actor: Actor, raw: dict[str, Any], _at: datetime | None
) -> Outcome:
    values = _values(DEVICES, raw)
    if "asset_name" in values:
        asset_name = values.pop("asset_name")
        asset = service.find_asset_by_name(session, asset_name) if asset_name else None
        if asset_name and asset is None:
            raise InvalidDataError(
                f"Техника {asset_name} не найдена: добавьте её на лист «Техника»"
            )
        values["asset_id"] = asset.id if asset else None
    device = service.find_device(session, name=values["name"], imei=values.get("imei"))
    if device is None:
        service.create_device(session, actor, DeviceCreate.model_validate(values))
        return "created"
    _, changed = service.update_device(
        session, actor, device.id, DeviceUpdate.model_validate(values)
    )
    return _outcome(changed)


Handler = Callable[[Session, Actor, dict[str, Any], datetime | None], Outcome]
_HANDLERS: dict[str, Handler] = {
    ASSETS.title: _import_asset,
    SITES.title: _import_site,
    CELLS.title: _import_cell,
    DEVICES.title: _import_device,
}


def import_workbook(
    session: Session, actor: Actor, data: bytes, effective_at: datetime | None
) -> list[SheetReport]:
    """Apply the workbook within the caller's transaction and report per sheet.

    The caller commits only if there are no errors (and it is not a dry run).
    """
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a variety of errors for broken files
        raise InvalidDataError("Не удалось прочитать файл: ожидается .xlsx") from exc

    if not any(spec.title in wb.sheetnames for spec in IMPORT_ORDER):
        wb.close()
        sheets = ", ".join(f"«{spec.title}»" for spec in IMPORT_ORDER)
        raise InvalidDataError(
            f"Это не шаблон импорта: в файле нет листов {sheets}. Скачайте шаблон на этой странице "
            "и перенесите данные в него. Файл Site Data оператора сначала переведите в шаблон "
            "командой dtat convert-kcell"
        )
    if effective_at is not None:
        effective_at = service.effective_time(effective_at)

    reports = []
    for spec in IMPORT_ORDER:
        report = SheetReport(sheet=spec.title)
        reports.append(report)
        if spec.title not in wb.sheetnames:
            report.present = False
            continue
        handler = _HANDLERS[spec.title]
        try:
            rows = list(_iter_rows(wb[spec.title], spec))
        except InvalidDataError as exc:
            report.errors.append(RowError(row=1, message=exc.message))
            continue
        for row_num, raw in rows:
            try:
                with session.begin_nested():
                    outcome = handler(session, actor, raw, effective_at)
            except ValidationError as exc:
                report.errors.append(RowError(row=row_num, message=_pydantic_message(spec, exc)))
                continue
            except DomainError as exc:
                report.errors.append(RowError(row=row_num, message=exc.message))
                continue
            setattr(report, outcome, getattr(report, outcome) + 1)
    wb.close()
    return reports
