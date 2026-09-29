"""Convert an operator's "Site Data" workbook (Kcell layout) into the inventory import template.

    dtat convert-kcell Site_Data_Kcell.xlsx inventory.xlsx --position 34348=46.949468,79.938039

Mapping (source columns → template):
- SITENAME "ERBS_34366_AKTOGCAREER_1_KK" → site code "34366", site name "AKTOGCAREER_1_KK",
  eNB ID 34366 (assumed: the number in the name; check against a phone or CM export), eNB name;
- EUTRANCELL → cell name, Cell ID, PHYSICALCELLID, EARFCNDL (UL = DL + 18000 for bands 1–28),
  AZIMUT, HEIGHT, tilts, antenna type;
- indoor/DAS cells get no azimuth (drawn as a circle); their type goes to the cell notes;
- remote sectors ("Remote сектор ERBS_<code>[_<name>]" in the comment, or other coordinates than the
  rest of the site) become their own sites, with "Код сайта eNB" pointing back to the eNodeB's site.
  The source gives remote sectors the eNodeB's coordinates: pass the real ones with `positions`.
"""

import re
from dataclasses import dataclass, field
from io import BytesIO
from typing import Any, BinaryIO

from openpyxl import load_workbook

from dtat.inventory.excel import CELLS, SITES, build_workbook

OUTDOOR_BEAMWIDTH_DEG = 65.0  # horizontal beamwidth of typical macro panels (HW ADU, Kathrein)
UL_OFFSET = 18000  # EARFCN UL − EARFCN DL for FDD bands 1–28
REMOTE_RE = re.compile(r"Remote сектор\s+ERBS_([A-Za-z0-9]+)(?:_(\S+))?", re.IGNORECASE)
SITENAME_RE = re.compile(r"^ERBS_(\d+)_(.+)$")
REQUIRED = ("SITENAME", "EUTRANCELL", "Cell ID", "Site type", "EARFCNDL", "PHYSICALCELLID")


@dataclass
class SiteRow:
    code: str
    name: str
    lat: float
    lon: float
    notes: list[str] = field(default_factory=list)


@dataclass
class Conversion:
    sites: list[SiteRow]
    cells: list[dict[str, Any]]

    @property
    def remote_cells(self) -> int:
        return sum(1 for c in self.cells if c["enb_site_code"])


def _number(value: Any) -> float | None:
    if value is None or (isinstance(value, str) and value.strip() in ("", "-")):
        return None
    return float(value)


def _text(value: Any) -> str:
    return " ".join(str(value or "").split())


def convert(
    source: BinaryIO | str, positions: dict[str, tuple[float, float]] | None = None
) -> Conversion:
    positions = positions or {}
    ws = load_workbook(source, data_only=True).active
    assert ws is not None
    header = [_text(c) for c in next(ws.iter_rows(max_row=1, values_only=True))]
    missing = [name for name in REQUIRED if name not in header]
    if missing:
        raise ValueError(f"Нет столбцов: {', '.join(missing)}")
    col = {name: header.index(name) for name in header if name}
    # Unnamed column at the end: free-text comment (remote sectors, tower, building).
    comment_idx = len(header) - 1

    def get(row: tuple[Any, ...], name: str) -> Any:
        return row[col[name]] if name in col else None

    sites: dict[str, SiteRow] = {}
    source_coords: dict[str, tuple[float, float]] = {}  # eNodeB site → coordinates in the source
    located: dict[tuple[float, float], str] = {}  # source coordinates → site code
    cells: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        sitename = _text(get(row, "SITENAME"))
        if not sitename:
            continue
        match = SITENAME_RE.match(sitename)
        if not match:
            raise ValueError(f"SITENAME не в формате ERBS_<номер>_<имя>: {sitename}")
        enb_site_code, site_name = match.groups()
        lat, lon = float(get(row, "LATITUDE")), float(get(row, "LONGITUDE"))
        comment = _text(row[comment_idx]) if len(row) > comment_idx else ""
        if enb_site_code not in source_coords:  # first row of this eNodeB
            # The site may already exist as another eNodeB's remote location.
            home = sites.setdefault(enb_site_code, SiteRow(enb_site_code, site_name, lat, lon))
            home.name = site_name
            home.lat, home.lon = positions.get(enb_site_code, (lat, lon))
            flags = [f"{n}={get(row, n)}" for n in ("Tower", "ДГУ") if get(row, n)]
            if flags:
                home.notes.append("Kcell: " + ", ".join(flags))
            source_coords[enb_site_code] = (lat, lon)
            located.setdefault((lat, lon), enb_site_code)
        home = sites[enb_site_code]

        # Where the antenna is: the eNodeB's site, or a remote location.
        location = home
        remote = REMOTE_RE.search(comment)
        if remote or (lat, lon) != source_coords[enb_site_code]:
            code = remote.group(1) if remote else None
            if code is None:  # same coordinates as an already known site
                code = located.get((lat, lon))
            if code is None:
                raise ValueError(f"Выносной сектор без имени площадки: {get(row, 'EUTRANCELL')}")
            if code not in sites:
                name = remote.group(2) if remote and remote.group(2) else code
                sites[code] = SiteRow(
                    code,
                    name,
                    *positions.get(code, (lat, lon)),
                    [f"выносные секторы eNodeB {enb_site_code}"],
                )
                if (lat, lon) != source_coords[enb_site_code]:
                    located.setdefault((lat, lon), code)
            location = sites[code]
            comment = "" if remote else comment
        elif comment and "вышк" in comment.lower():
            home.notes.append(comment)
            comment = ""

        site_type = _text(get(row, "Site type"))
        indoor = not site_type.lower().startswith("outdoor")
        antennas = _number(get(row, "Number of antennas"))
        notes = [
            site_type if site_type.lower() != "outdoor" else "",
            f"антенн: {int(antennas)}" if antennas and antennas > 1 else "",
            comment,
        ]
        earfcn = int(get(row, "EARFCNDL"))
        cells.append(
            {
                "site_code": location.code,
                "enb_id": int(enb_site_code),
                "enb_name": sitename,
                "enb_site_code": enb_site_code if location is not home else None,
                "local_cell_id": int(get(row, "Cell ID")),
                "name": _text(get(row, "EUTRANCELL")),
                "status": "в работе",
                "pci": int(get(row, "PHYSICALCELLID")),
                "earfcn_dl": earfcn,
                "earfcn_ul": earfcn + UL_OFFSET,
                "antenna_model": _text(get(row, "Antenna type")) or None,
                "height_m": _number(get(row, "HEIGHT")),
                "azimuth_deg": None if indoor else _number(get(row, "AZIMUT")),
                "mech_tilt_deg": _number(get(row, "Mechanical tilt")),
                "elec_tilt_deg": _number(get(row, "Electrical tilt")),
                "beamwidth_deg": None if indoor else OUTDOOR_BEAMWIDTH_DEG,
                "notes": "; ".join(n for n in notes if n) or None,
            }
        )
    return Conversion(list(sites.values()), cells)


def write_template(conversion: Conversion) -> bytes:
    wb = load_workbook(BytesIO(build_workbook(None)))
    site_ws, cell_ws = wb[SITES.title], wb[CELLS.title]
    for site in conversion.sites:
        values = {
            "code": site.code,
            "name": site.name,
            "kind": "стационарный",
            "status": "в работе",
            "lat": site.lat,
            "lon": site.lon,
            "notes": "; ".join(site.notes) or None,
        }
        site_ws.append([values.get(c.key) for c in SITES.columns])
    for cell in conversion.cells:
        cell_ws.append([cell.get(c.key) for c in CELLS.columns])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
