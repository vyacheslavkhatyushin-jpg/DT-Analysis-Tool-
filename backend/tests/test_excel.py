from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from dtat.inventory.excel import CELLS, SITES
from dtat.synthetic.quarry import seed_quarry

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _upload(client: TestClient, data: bytes, *, dry_run: bool, **form: str) -> dict:
    response = client.post(
        "/api/v1/inventory/import",
        files={"file": ("inventory.xlsx", data, XLSX)},
        data={"dry_run": str(dry_run).lower(), **form},
    )
    assert response.status_code == 200, response.text
    return response.json()


def _sheet(report: dict, title: str) -> dict:
    return next(s for s in report["sheets"] if s["sheet"] == title)


def _workbook_with(rows: dict[str, list[list[object]]]) -> bytes:
    template_wb = load_workbook(BytesIO(_template_bytes()))
    for title, sheet_rows in rows.items():
        ws = template_wb[title]
        for row in sheet_rows:
            ws.append(row)
    buffer = BytesIO()
    template_wb.save(buffer)
    return buffer.getvalue()


_TEMPLATE: bytes | None = None


def _template_bytes() -> bytes:
    from dtat.inventory.excel import build_workbook

    global _TEMPLATE
    if _TEMPLATE is None:
        _TEMPLATE = build_workbook(None)
    return _TEMPLATE


def test_template_has_all_sheets(viewer: TestClient) -> None:
    response = viewer.get("/api/v1/inventory/template.xlsx")
    assert response.status_code == 200
    wb = load_workbook(BytesIO(response.content))
    assert wb.sheetnames == ["Инструкция", "Сайты", "Соты", "Техника", "Устройства"]
    assert [c.value for c in wb["Сайты"][1]] == [c.header for c in SITES.columns]


def test_export_import_round_trip_is_unchanged(engineer: TestClient, session: Session) -> None:
    seed_quarry(session)
    session.commit()
    exported = engineer.get("/api/v1/inventory/export.xlsx").content
    report = _upload(engineer, exported, dry_run=True)
    assert report["total_errors"] == 0
    for sheet in report["sheets"]:
        assert sheet["created"] == sheet["updated"] == 0, sheet
        assert sheet["unchanged"] > 0


def test_import_creates_and_updates(engineer: TestClient) -> None:
    site_row = ["KR-01", "Борт север", "стационарный", "в работе", 54.1, 87.1]
    # Код сайта, eNB ID, Имя eNB, Код сайта eNB, Cell ID, ECI, Имя соты, Статус, PCI, EARFCN DL
    cell_row = ["KR-01", 170001, "KR-01", None, 1, None, "KR-01-1", "в работе", 5, 1300]
    data = _workbook_with({"Сайты": [site_row], "Соты": [cell_row]})

    preview = _upload(engineer, data, dry_run=True)
    assert _sheet(preview, "Сайты")["created"] == 1
    assert _sheet(preview, "Соты")["created"] == 1
    assert engineer.get("/api/v1/sites").json() == []  # dry run applies nothing

    applied = _upload(engineer, data, dry_run=False)
    assert applied["applied"] is True
    cell = engineer.get("/api/v1/cells").json()[0]
    assert cell["eci"] == 170001 * 256 + 1
    assert cell["pci"] == 5

    changed = _workbook_with({"Сайты": [site_row], "Соты": [[*cell_row[:8], 8, 1300]]})
    report = _upload(engineer, changed, dry_run=False, effective_at="2026-01-15T08:00:00Z")
    assert _sheet(report, "Соты")["updated"] == 1
    assert _sheet(report, "Сайты")["unchanged"] == 1
    versions = engineer.get(f"/api/v1/cells/{cell['id']}/versions").json()
    assert versions[0]["pci"] == 8
    assert versions[0]["valid_from"].startswith("2026-01-15T08:00:00")
    sources = {c["source"] for c in engineer.get("/api/v1/changes").json()}
    assert sources == {"import"}


def test_import_with_errors_applies_nothing(engineer: TestClient) -> None:
    data = _workbook_with(
        {
            "Сайты": [
                ["OK-1", None, None, None, 54.0, 87.0],
                ["BAD-1", None, "шалаш", None, 54.0, 87.0],
                ["BAD-2", None, None, None, "север", 87.0],
            ],
            "Соты": [["NOPE", 1, None, None, 1, None, None, None, 1, 1300]],
        }
    )
    report = _upload(engineer, data, dry_run=False)
    assert report["applied"] is False
    errors = _sheet(report, "Сайты")["errors"]
    assert [e["row"] for e in errors] == [3, 4]
    assert "Тип площадки" in errors[0]["message"]
    assert "Широта" in errors[1]["message"]
    assert "NOPE" in _sheet(report, "Соты")["errors"][0]["message"]
    assert engineer.get("/api/v1/sites").json() == []


def test_missing_required_column(engineer: TestClient) -> None:
    template = load_workbook(BytesIO(_template_bytes()))
    ws = template["Соты"]
    ws.delete_cols([c.key for c in CELLS.columns].index("pci") + 1)
    buffer = BytesIO()
    template.save(buffer)
    report = _upload(engineer, buffer.getvalue(), dry_run=True)
    assert "PCI" in _sheet(report, "Соты")["errors"][0]["message"]


def test_devices_link_to_assets(engineer: TestClient) -> None:
    data = _workbook_with(
        {
            "Техника": [["С-101", "самосвал", "130 т"]],
            "Устройства": [
                ["RTR-С-101", "роутер", None, "490154203237518", None, None, None, "С-101"]
            ],
        }
    )
    report = _upload(engineer, data, dry_run=False)
    assert report["applied"], report
    device = engineer.get("/api/v1/devices").json()[0]
    asset = engineer.get("/api/v1/assets").json()[0]
    assert device["asset_id"] == asset["id"]


def test_not_an_xlsx(engineer: TestClient) -> None:
    response = engineer.post(
        "/api/v1/inventory/import",
        files={"file": ("x.xlsx", b"not a workbook", XLSX)},
        data={"dry_run": "true"},
    )
    assert response.status_code == 422


def test_workbook_without_template_sheets(engineer: TestClient) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["SITENAME", "EUTRANCELL", "Cell ID"])
    buffer = BytesIO()
    wb.save(buffer)
    response = engineer.post(
        "/api/v1/inventory/import",
        files={"file": ("site_data.xlsx", buffer.getvalue(), XLSX)},
        data={"dry_run": "true"},
    )
    assert response.status_code == 422
    assert "не шаблон импорта" in response.json()["detail"]


def test_reimport_with_corrected_enb_id_renumbers(engineer: TestClient) -> None:
    sites: list[list[object]] = [["34344", "Sulphide", None, None, 46.947, 79.935]]
    cells: list[list[object]] = [
        ["34344", enb, "ERBS_34344", None, local, None, f"P_534344-{local}", None, pci, 6200]
        for local, pci in ((70, 43), (72, 40))
        for enb in (34344,)
    ]
    assert _upload(engineer, _workbook_with({"Сайты": sites, "Соты": cells}), dry_run=False)[
        "applied"
    ]
    # The phones report eNB 534344: the same cells under the corrected eNB ID.
    fixed: list[list[object]] = [[*row[:1], 534344, *row[2:]] for row in cells]
    report = _upload(engineer, _workbook_with({"Сайты": sites, "Соты": fixed}), dry_run=False)
    # The first row renumbers the eNodeB; the second finds it under the new ID.
    assert report["applied"] and _sheet(report, "Соты")["updated"] == 1, report
    enodebs = engineer.get("/api/v1/enodebs").json()
    assert [e["enb_id"] for e in enodebs] == [534344]
    by_name = {c["name"]: c for c in engineer.get("/api/v1/cells").json()}
    assert by_name["P_534344-70"]["eci"] == 534344 * 256 + 70


def test_import_remote_sector(engineer: TestClient) -> None:
    sites: list[list[object]] = [
        ["ZH", "Жанар", None, None, 46.894, 79.612],
        ["ZH-PS3", "PS3", None, None, 46.958, 79.858],
    ]
    cells: list[list[object]] = [
        ["ZH", 34371, "ERBS_34371", None, 11, None, "ZH-A", None, 7, 6200],
        ["ZH-PS3", 34371, "ERBS_34371", "ZH", 14, None, "ZH-D", None, 17, 6200],
    ]
    report = _upload(engineer, _workbook_with({"Сайты": sites, "Соты": cells}), dry_run=False)
    assert report["applied"], report
    by_name = {c["name"]: c for c in engineer.get("/api/v1/cells").json()}
    assert by_name["ZH-A"]["site_code"] == "ZH"
    assert by_name["ZH-D"]["site_code"] == "ZH-PS3"
    assert by_name["ZH-D"]["enodeb_site_code"] == "ZH"

    exported = engineer.get("/api/v1/inventory/export.xlsx").content
    ws = load_workbook(BytesIO(exported))["Соты"]
    headers = [c.value for c in ws[1]]
    rows = {r[headers.index("Имя соты")]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert rows["ZH-D"][headers.index("Код сайта eNB")] == "ZH"
    assert rows["ZH-A"][headers.index("Код сайта eNB")] is None
    round_trip = _upload(engineer, exported, dry_run=True)
    assert all(s["created"] == s["updated"] == 0 for s in round_trip["sheets"]), round_trip
