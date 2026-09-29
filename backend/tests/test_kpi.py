from datetime import date
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook

from tests.helpers import create_cell, create_site

HEADER = [
    None, "Date", "Hour", "ERBS Name", "EUtranCell Id", "RRC Setup Success Rate, % ",
    "ERAB Drop Rate,% wo UE lost", "Mobility Success Rate,%", "DL Payload,GB", "UL Payload,GB",
    "Some new KPI",
]  # fmt: skip
DAY = date(2026, 8, 29)
ENB = "ERBS_34340_SOUTH_1_KP"


def _report(rows: list[list[object]]) -> bytes:
    """The operator's layout: a title row, the header on row 2, an eNodeB sheet without cells."""
    wb = Workbook()
    by_bs = wb.active
    assert by_bs is not None
    by_bs.title = "ALL KPI by BS"
    by_bs.append([None, "Date", "Hour", "ERBS Name", "ERAB Drop Rate,%"])
    by_cell = wb.create_sheet("ALL KPI by Cell")
    by_cell.append([])
    by_cell.append(HEADER)
    for row in rows:
        by_cell.append([None, DAY, *row])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


ROWS: list[list[object]] = [
    # Hour, eNB, cell, RRC SR, drop, HO SR, DL GB, UL GB, unknown
    [0, ENB, "C-70", 100, "#DIV/0", 90, 1, 1, 5],
    [1, ENB, "C-70", 96, 4, 100, 3, 1, 5],
    [0, ENB, "C-72", 99.5, 0, 99, 0.5, 0.5, 5],
    [0, ENB, "X-74", 99.9, 0, 99, 0.1, 0.1, 5],
]


def _import(client: TestClient, data: bytes, **form: str) -> dict:
    response = client.post(
        "/api/v1/kpi/import",
        files={"file": ("kpi.xlsx", data, "application/octet-stream")},
        data=form,
    )
    assert response.status_code == 200, response.text
    return response.json()


def _inventory(client: TestClient) -> dict[str, int]:
    site = create_site(client)
    response = client.post(
        "/api/v1/enodebs",
        json={"site_id": site["id"], "enb_id": 34340, "name": "ERBS_34340_SOUTH_KP"},
    )
    assert response.status_code == 201, response.text
    enodeb = response.json()["id"]
    cells = {}
    for local_id, name in ((70, "C-70"), (72, "C-72"), (74, "OLD-74")):
        cells[name] = create_cell(client, enodeb, local_cell_id=local_id, name=name)["id"]
    return cells


def test_import_links_by_name_and_summarises(engineer: TestClient) -> None:
    cells = _inventory(engineer)
    result = _import(engineer, _report(ROWS))
    assert result["sheet"] == "ALL KPI by Cell"
    assert result["unknown_columns"] == ["Some new KPI"]
    assert result["new_cells"] == ["C-70", "C-72", "X-74"]
    assert result["unlinked_cells"] == ["X-74"]
    # Local time of the report (Asia/Almaty, UTC+5 by default).
    assert result["file"]["period_start"] == "2026-08-28T19:00:00Z"
    assert result["file"]["period_end"] == "2026-08-28T21:00:00Z"
    assert result["file"]["samples"] == 4 * 5 - 1  # "#DIV/0" is no value

    summary = engineer.get(
        "/api/v1/kpi/summary",
        params={"start": "2026-08-28T00:00:00Z", "end": "2026-08-30T00:00:00Z"},
    ).json()
    c70 = next(c for c in summary["cells"] if c["cell_name"] == "C-70")
    assert c70["cell_id"] == cells["C-70"]
    rrc, ho, drop = (
        c70["values"]["rrc_sr"],
        c70["values"]["ho_sr"],
        c70["values"]["erab_drop_wo_ue_lost"],
    )
    assert (rrc["value"], rrc["level"], rrc["worst"], rrc["worst_level"]) == (98, "warn", 96, "bad")
    # Weighted by the hour's traffic: (90 × 2 GB + 100 × 4 GB) / 6 GB.
    assert round(ho["value"], 2) == 96.67 and ho["level"] == "warn"
    assert (drop["value"], drop["hours"], drop["level"]) == (4, 1, "bad")
    assert c70["values"]["dl_volume"]["value"] == 4 and c70["values"]["dl_volume"]["level"] is None

    series = engineer.get(f"/api/v1/cells/{cells['C-70']}/kpi", params={"kpi": ["rrc_sr"]}).json()
    assert [p["values"] for p in series["points"]] == [{"rrc_sr": 100}, {"rrc_sr": 96}]


def test_reimport_replaces_values(engineer: TestClient) -> None:
    _import(engineer, _report(ROWS))
    changed = [[*ROWS[0][:3], 50, *ROWS[0][4:]], *ROWS[1:]]
    result = _import(engineer, _report(changed))
    assert result["new_cells"] == []
    summary = engineer.get("/api/v1/kpi/summary").json()
    c70 = next(c for c in summary["cells"] if c["cell_name"] == "C-70")
    assert c70["values"]["rrc_sr"]["value"] == 73  # (50 + 96) / 2, the old 100 is gone
    assert len(engineer.get("/api/v1/kpi/files").json()) == 2


def test_reconciliation_and_link_with_rename(engineer: TestClient) -> None:
    cells = _inventory(engineer)
    _import(engineer, _report(ROWS))
    report = engineer.get("/api/v1/kpi/reconciliation").json()
    assert [u["cell_name"] for u in report["unlinked"]] == ["X-74"]
    assert [c["name"] for c in report["unlinked"][0]["candidates"]] == ["OLD-74"]
    assert [c["name"] for c in report["cells_without_kpi"]] == ["OLD-74"]
    assert report["enb_names"] == [
        {
            "enodeb_id": report["enb_names"][0]["enodeb_id"],
            "enb_id": 34340,
            "name": "ERBS_34340_SOUTH_KP",
            "kpi_name": ENB,
        }
    ]

    kpi_cell = report["unlinked"][0]["kpi_cell_id"]
    response = engineer.put(
        f"/api/v1/kpi/operator-cells/{kpi_cell}/link", json={"cell_id": cells["OLD-74"]}
    )
    assert response.status_code == 200, response.text
    assert engineer.get(f"/api/v1/cells/{cells['OLD-74']}").json()["name"] == "X-74"
    report = engineer.get("/api/v1/kpi/reconciliation").json()
    assert report["unlinked"] == [] and report["cells_without_kpi"] == []
    series = engineer.get(f"/api/v1/cells/{cells['OLD-74']}/kpi").json()
    assert len(series["points"]) == 1


def test_link_rename_refuses_taken_name(engineer: TestClient) -> None:
    cells = _inventory(engineer)
    _import(engineer, _report(ROWS))
    kpi_cell = engineer.get("/api/v1/kpi/reconciliation").json()["unlinked"][0]["kpi_cell_id"]
    engineer.patch(f"/api/v1/cells/{cells['C-72']}", json={"name": "X-74"})
    response = engineer.put(
        f"/api/v1/kpi/operator-cells/{kpi_cell}/link", json={"cell_id": cells["OLD-74"]}
    )
    assert response.status_code == 422
    assert "занято" in response.json()["detail"]


def test_import_errors(engineer: TestClient) -> None:
    files = {"file": ("kpi.xlsx", _report(ROWS), "application/octet-stream")}
    response = engineer.post("/api/v1/kpi/import", files=files, data={"timezone": "Mars/Base"})
    assert response.status_code == 422 and response.json()["field"] == "timezone"

    wb = Workbook()
    buffer = BytesIO()
    wb.save(buffer)
    response = engineer.post(
        "/api/v1/kpi/import", files={"file": ("x.xlsx", buffer.getvalue(), "application/x")}
    )
    assert response.status_code == 422
    assert "Не найден лист" in response.json()["detail"]


def test_viewer_reads_but_cannot_import(viewer: TestClient) -> None:
    files = {"file": ("kpi.xlsx", _report(ROWS), "application/octet-stream")}
    assert viewer.post("/api/v1/kpi/import", files=files).status_code == 403

    summary = viewer.get("/api/v1/kpi/summary").json()
    assert summary["cells"] == [] and summary["start"] is None
    catalogue = viewer.get("/api/v1/kpi/catalogue").json()
    assert {"code": "ho_sr", "better": "high", "warn": 98.0, "bad": 95.0}.items() <= next(
        k for k in catalogue if k["code"] == "ho_sr"
    ).items()
